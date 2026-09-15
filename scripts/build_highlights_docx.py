"""Build the Highlights + eTOC blurb as a single Word document
(Cell Press requires this as one .docx, separate from the main text).
Source of truth: docs/highlights_and_etoc.md
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

R = Path(__file__).resolve().parents[1]
SRC = R / 'docs' / 'highlights_and_etoc.md'
OUT = R / 'docs' / 'highlights_and_etoc.docx'

t = SRC.read_text(encoding='utf-8')
highlights = re.findall(r'^\d+\.\s+(.+)$', t.split('## eTOC')[0], re.M)
after = re.sub(r'^[^\n]*\n', '', t.split('## eTOC')[1], count=1)  # drop heading remainder
etoc = re.sub(r'\s+', ' ', after).strip()

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

doc.add_heading('Highlights', level=1)
for h in highlights:
    p = doc.add_paragraph(h, style='List Bullet')
    p.paragraph_format.space_after = Pt(6)

doc.add_heading('eTOC blurb', level=1)
p = doc.add_paragraph(etoc)
p.paragraph_format.space_after = Pt(6)

doc.save(OUT)
print('written:', OUT)
print('highlights:', len(highlights), '| chars:', [len(x) for x in highlights])
print('etoc words:', len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'\-]*", etoc)))
