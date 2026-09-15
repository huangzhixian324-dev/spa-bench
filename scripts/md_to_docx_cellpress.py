"""Markdown -> Word (docx) for Cell Press submission.
Cell Press wants: single editable main-text document (title through references
and figure legends), figures uploaded separately. Tables inline. In-text
citations rendered as superscripts (Cell Press style).
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = Path(__file__).resolve().parents[1]
md = (REPO / 'docs/manuscript_v34_nar.md').read_text(encoding='utf-8')

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)

# title
m = re.search(r'^# (.+)$', md, re.M)
if m:
    p = doc.add_paragraph()
    r = p.add_run(m.group(1))
    r.bold = True
    r.font.size = Pt(16)
    md = md.replace(m.group(0), '', 1)

# split into blocks
lines = md.splitlines()
i = 0
buf = []


def flush_para():
    global buf
    if not buf:
        return
    text = ' '.join(buf).strip()
    buf = []
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # 处理行内标记：**bold**、*italic*、[n] 上标
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\[\d+(?:[,\-]\s?\d+)*\])', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        elif re.fullmatch(r'\[\d+(?:[,\-]\s?\d+)*\]', part):
            r = p.add_run(part); r.font.superscript = True
        else:
            p.add_run(part)


while i < len(lines):
    line = lines[i]
    ls = line.strip()
    if ls.startswith('## '):
        flush_para()
        doc.add_heading(ls[3:], level=1)
    elif ls.startswith('### '):
        flush_para()
        doc.add_heading(ls[4:], level=2)
    elif ls.startswith('|'):
        flush_para()
        # collect table
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
                rows.append(cells)
            i += 1
        if rows:
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    if ci < len(t.rows[ri].cells):
                        t.rows[ri].cells[ci].text = cell
                        for r_ in t.rows[ri].cells[ci].paragraphs[0].runs:
                            r_.font.size = Pt(8)
        continue
    elif ls.startswith('- ') or ls.startswith('* ') and not ls.startswith('**'):
        flush_para()
        p = doc.add_paragraph(ls[2:], style='List Bullet')
    elif re.match(r'^\d+\.\s', ls):
        flush_para()
        p = doc.add_paragraph(re.sub(r'^\d+\.\s', '', ls), style='List Number')
    elif ls == '' or ls == '---':
        flush_para()
    else:
        buf.append(ls)
    i += 1
flush_para()

out = REPO / 'docs/SPATBench_CellSystems_manuscript_v35.docx'
doc.save(out)
print('written:', out, out.stat().st_size // 1024, 'KB')
